from openpyxl import load_workbook
from pathlib import Path
import re

xlsx = Path('/Users/ketangchen/Documents/000_20260114dev/research_laboratory/ChatAFL-master/ASE-Paper-Verified-LLM/conference.xlsx')
bib_path = Path('/Users/ketangchen/Documents/000_20260114dev/research_laboratory/ChatAFL-master/ASE-Paper-Verified-LLM/references.bib')
wb = load_workbook(xlsx, data_only=True)
bib_text = bib_path.read_text().lower()
bib_norm = re.sub(r'[^a-z0-9]+', '', bib_text)
patterns = [
    'protocol', 'stateful', 'network', 'llm', 'large language', 'grammar',
    'document', 'rfc', 'message sequence', 'fuzz driver', 'service-aware',
    'request', 'response', 'smtp', 'ftp', 'mqtt', 'rtsp'
]
rows_out = []
for ws in wb.worksheets:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        continue
    header = [str(x).strip() if x is not None else '' for x in rows[0]]
    for row in rows[1:]:
        vals = {header[i]: row[i] if i < len(row) else None for i in range(len(header))}
        title = str(vals.get('论文') or '').strip()
        if not title:
            continue
        keywords = str(vals.get('关键字分类问题') or '').strip()
        note = str(vals.get('老师建议') or '').strip()
        blob = ' '.join([title.lower(), keywords.lower(), note.lower()])
        if not any(p in blob for p in patterns):
            continue
        title_norm = re.sub(r'[^a-z0-9]+', '', title.lower())
        if title_norm and title_norm in bib_norm:
            continue
        rows_out.append({
            'sheet': ws.title,
            'year': vals.get('年份'),
            'title': title,
            'authors': str(vals.get('作者') or '').strip(),
            'keywords': keywords,
            'advice': note,
            'link': str(vals.get('link') or '').strip(),
            'code': str(vals.get('代码链接') or '').strip(),
        })

for item in rows_out:
    print(f"{item['sheet']}\t{item['year']}\t{item['title']}\t{item['authors']}\t{item['keywords']}")
print(f'TOTAL\t{len(rows_out)}')
